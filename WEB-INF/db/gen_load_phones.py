"""Build WEB-INF/db/load_phones.sql from Phonedata.xlsx. Run once when the sheet changes."""
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

SRC = Path(r"c:\Users\mvplo\OneDrive\Documents\Phonedata.xlsx")
OUT = Path(__file__).with_name("load_phones.sql")


def digits(v):
    return "".join(c for c in str(v or "") if c.isdigit())


def fmt_phone(v):
    d = digits(v)
    if len(d) == 10:
        return "%s-%s-%s" % (d[:3], d[3:6], d[6:])
    return d


def sql_str(v):
    if v is None:
        return "NULL"
    s = str(v).strip()
    if s == "" or s.lower() == "none":
        return "NULL"
    return "'" + s.replace("\\", "\\\\").replace("'", "''") + "'"


def sql_dt(v):
    if v is None:
        return "NULL"
    if isinstance(v, datetime):
        return "'" + v.strftime("%Y-%m-%d %H:%M:%S") + "'"
    if isinstance(v, date):
        return "'" + v.strftime("%Y-%m-%d") + " 00:00:00'"
    s = str(v).strip()
    if not s:
        return "NULL"
    for fmt in ("%Y-%m-%d %H:%M:%S", "%Y-%m-%d", "%m/%d/%Y"):
        try:
            return "'" + datetime.strptime(s[:19], fmt).strftime("%Y-%m-%d %H:%M:%S") + "'"
        except Exception:
            pass
    return "NULL"


def main():
    wb = load_workbook(SRC, data_only=True)
    ws = wb["Sheet2"]
    rows = list(ws.iter_rows(values_only=True))
    hdr = [str(c).strip() if c else "" for c in rows[0]]
    data = []
    seen = set()
    for r in rows[1:]:
        if not any(c is not None and str(c).strip() for c in r):
            continue
        d = dict(zip(hdr, r))
        p = digits(d.get("Phone Number"))
        if not p or p in seen:
            continue
        seen.add(p)
        data.append(d)

    lines = [
        "-- Load cleaned Phonedata.xlsx into fleetdb.phones",
        "-- Match on digits-only phone number. Does not hard-delete existing rows.",
        "-- Excel Suspend -> PHONESTATUS=4 (Inactive). Inserts default CURRENTSTATUS=0 (In Use).",
        "-- Run after alter_phones.sql. Review, then COMMIT;",
        "",
        "USE fleetdb;",
        "",
        "SET SQL_SAFE_UPDATES = 0;",
        "START TRANSACTION;",
        "",
        "-- %d spreadsheet rows" % len(data),
        "",
    ]

    for d in data:
        raw = digits(d.get("Phone Number"))
        phone = fmt_phone(d.get("Phone Number"))
        st = str(d.get("Phone Status") or "").strip().lower()
        pstatus = 4 if st in ("suspend", "suspended", "inactive") else 0
        notes_sql = sql_str(d.get("Additional Notes"))
        remark_set = "" if notes_sql == "NULL" else (", REMARKS = %s" % notes_sql)
        lines.extend([
            "UPDATE phones SET",
            "  PHONENUMBER = %s," % sql_str(phone),
            "  PHONESTATUS = %d," % pstatus,
            "  SERIALNUMBER = %s," % sql_str(d.get("Device In Use IMEI 1")),
            "  DEVICEMAKE = %s," % sql_str(d.get("Device In Use Make")),
            "  DEVICEMODEL = %s," % sql_str(d.get("Device In Use Model")),
            "  IMEI2 = %s," % sql_str(d.get("Device In Use IMEI 2")),
            "  IMSI = %s," % sql_str(d.get("Device In Use IMSI")),
            "  ICCID = %s," % sql_str(d.get("Device In Use ICCID")),
            "  EID = %s," % sql_str(d.get("Device In Use EID")),
            "  CONTRACTENDDATE = %s," % sql_dt(d.get("Contract End Date")),
            "  CONTRACTSTARTDATE = %s," % sql_dt(d.get("Contract Start Date")),
            "  DEVICEORDEREDDATE = %s," % sql_dt(d.get("Device Ordered Date")),
            "  DEVICEORDEREDIMEI = %s," % sql_str(d.get("Device Ordered IMEI")),
            "  DEVICEINUSEDATE = %s," % sql_dt(d.get("Device In Use Date")),
            "  UPDATE_USER = 'phone-load',",
            "  UPDATE_DATE = NOW()%s" % remark_set,
            "WHERE STATUS != 1 AND REPLACE(REPLACE(IFNULL(PHONENUMBER,''),'-',''),' ','') = '%s';" % raw,
            "",
            "INSERT INTO phones (ENTITYID, PHONENUMBER, PHONESTATUS, CURRENTSTATUS, SERIALNUMBER,",
            "  DEVICEMAKE, DEVICEMODEL, IMEI2, IMSI, ICCID, EID, REMARKS,",
            "  CONTRACTENDDATE, CONTRACTSTARTDATE, DEVICEORDEREDDATE, DEVICEORDEREDIMEI,",
            "  DEVICEINUSEDATE, CREATE_USER, CREATE_DATE, STATUS)",
            "SELECT 1,",
            "  %s, %d, 0, %s," % (sql_str(phone), pstatus, sql_str(d.get("Device In Use IMEI 1"))),
            "  %s, %s, %s," % (
                sql_str(d.get("Device In Use Make")),
                sql_str(d.get("Device In Use Model")),
                sql_str(d.get("Device In Use IMEI 2")),
            ),
            "  %s, %s, %s," % (
                sql_str(d.get("Device In Use IMSI")),
                sql_str(d.get("Device In Use ICCID")),
                sql_str(d.get("Device In Use EID")),
            ),
            "  %s, %s, %s," % (
                notes_sql,
                sql_dt(d.get("Contract End Date")),
                sql_dt(d.get("Contract Start Date")),
            ),
            "  %s, %s, %s," % (
                sql_dt(d.get("Device Ordered Date")),
                sql_str(d.get("Device Ordered IMEI")),
                sql_dt(d.get("Device In Use Date")),
            ),
            "  'phone-load', NOW(), 0",
            "FROM DUAL",
            "WHERE NOT EXISTS (SELECT 1 FROM phones z WHERE z.STATUS != 1 AND REPLACE(REPLACE(IFNULL(z.PHONENUMBER,''),'-',''),' ','') = '%s');" % raw,
            "",
        ])

    lines.extend([
        "SELECT 'phones after load' AS what, COUNT(*) n, SUM(PHONESTATUS=0) active, SUM(PHONESTATUS=4) inactive FROM phones WHERE STATUS != 1;",
        "",
        "SET SQL_SAFE_UPDATES = 1;",
        "-- COMMIT;",
        "-- ROLLBACK;",
        "",
    ])
    OUT.write_text("\n".join(lines), encoding="utf-8")
    print("wrote %s rows=%d bytes=%d" % (OUT, len(data), OUT.stat().st_size))


if __name__ == "__main__":
    main()
