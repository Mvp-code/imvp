#!/usr/bin/env python3
"""
Import MVP Fleet Inventory Report into fleetdb.

Sources (Excel):
  - Current Issues (Full Inventory)  -> vehicle upsert by VIN
  - Vehicle History                  -> vehicle_maintenance_log inserts

Safety:
  - Tags all new maintenance rows with CREATE_USER = BATCH_TAG
  - Tags newly inserted vehicles with CREATE_USER = BATCH_TAG
  - Writes touched VINs to fleet_import_batch / fleet_import_touch
  - Creates backup tables before mutating live data

Usage:
  python import_fleet_inventory.py
  python import_fleet_inventory.py --excel "C:\\path\\to\\report.xlsx" --dry-run
"""

from __future__ import annotations

import argparse
import hashlib
import re
import sys
from datetime import datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    import subprocess

    subprocess.check_call([sys.executable, "-m", "pip", "install", "openpyxl", "-q"])
    import openpyxl

try:
    import mysql.connector
except ImportError:
    import subprocess

    subprocess.check_call([sys.executable, "-m", "pip", "install", "mysql-connector-python", "-q"])
    import mysql.connector

DEFAULT_EXCEL = Path(r"C:\Users\mvplo\Downloads\MVP_Fleet_Inventory_Report_3.xlsx")
BATCH_TAG = "FLEET_INV_" + datetime.now().strftime("%Y%m%d_%H%M%S")
ENTITY_ID = 1
STATION = "DNK7"

# Excel Activity Type -> MAINT_CODE
ACTIVITY_MAP = [
    (re.compile(r"oil\s*change", re.I), "OIL_CHANGE"),
    (re.compile(r"tire\s*rotat|tire\s*service|goodyear|flat\s*tire", re.I), "TIRE_SERVICE"),
    (re.compile(r"brake", re.I), "BRAKE_SERVICE"),
    (re.compile(r"recall", re.I), "RECALL_SERVICE"),
    (re.compile(r"\btow\b", re.I), "TOW"),
    (re.compile(r"windshield|glass", re.I), "WINDSHIELD_GLASS"),
    (re.compile(r"dealer|liccardi|ciocca|fullerton|maaco|ray catena", re.I), "DEALER_REPAIR"),
    (re.compile(r"repair|case\s+\d+", re.I), "ONSITE_REPAIR"),
]


def map_activity(activity: str) -> str:
    a = (activity or "").strip()
    for rx, code in ACTIVITY_MAP:
        if rx.search(a):
            return code
    return "ONSITE_REPAIR"


def sql_str(v) -> str:
    if v is None:
        return "NULL"
    s = str(v).strip()
    if not s or s.lower() in ("none", "null", "n/a", "no data on file"):
        return "NULL"
    s = s.replace("\\", "\\\\").replace("'", "''")
    return f"'{s}'"


def sql_date(v) -> str:
    if v is None:
        return "NULL"
    if isinstance(v, datetime):
        return f"'{v.strftime('%Y-%m-%d')}'"
    s = str(v).strip()
    if not s or s.lower() in ("none", "null", "undated", "n/a"):
        return "NULL"
    # already YYYY-MM-DD-ish
    m = re.match(r"^(\d{4}-\d{2}-\d{2})", s)
    if m:
        return f"'{m.group(1)}'"
    m = re.match(r"^(\d{1,2})/(\d{1,2})/(\d{4})", s)
    if m:
        return f"'{int(m.group(3)):04d}-{int(m.group(1)):02d}-{int(m.group(2)):02d}'"
    return "NULL"


def truthy(v) -> bool:
    if v is None:
        return False
    if isinstance(v, bool):
        return v
    s = str(v).strip().lower()
    return s not in ("", "0", "false", "no", "none", "n/a", "null")


def normalize_van(van: str) -> str:
    v = (van or "").strip()
    if not v or v.lower() in ("(not in roster)", "(unnamed)"):
        return ""
    return v


def ownership_to_vehicle_type(ownership: str) -> str:
    o = (ownership or "").strip().upper()
    if o == "RENTAL":
        return "Rental"
    if "AMAZON" in o:
        return "Amazon-Owned"
    return ownership or ""


def op_status(excel_status: str) -> tuple[int, int]:
    """Return (OPERATIONALSTATUS, STATUS record)."""
    s = (excel_status or "").strip().upper()
    if s == "OPERATIONAL":
        return 0, 0
    if s == "GROUNDED":
        return 1, 0
    if s == "INACTIVE":
        return 1, 4  # RecordStatus.INACTIVE
    return 0, 0


def connect():
    return mysql.connector.connect(
        host="localhost",
        user="root",
        password="admin",
        database="fleetdb",
        autocommit=False,
    )


def ensure_batch_tables(cur):
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS fleet_import_batch (
          BATCH_TAG varchar(60) NOT NULL PRIMARY KEY,
          EXCEL_PATH varchar(500),
          FILE_SHA1 varchar(40),
          STARTED_AT datetime,
          FINISHED_AT datetime,
          VEHICLES_INSERTED int DEFAULT 0,
          VEHICLES_UPDATED int DEFAULT 0,
          HISTORY_INSERTED int DEFAULT 0,
          NOTES varchar(500)
        )
        """
    )
    cur.execute(
        """
        CREATE TABLE IF NOT EXISTS fleet_import_touch (
          TOUCHID int NOT NULL AUTO_INCREMENT PRIMARY KEY,
          BATCH_TAG varchar(60) NOT NULL,
          VINNUMBER varchar(50),
          VEHICLEID int,
          ACTION varchar(20),
          CREATE_DATE datetime,
          KEY idx_batch (BATCH_TAG),
          KEY idx_vin (VINNUMBER)
        )
        """
    )


def backup_tables(cur, batch_tag: str):
    safe = re.sub(r"[^A-Za-z0-9_]", "_", batch_tag.lower())
    veh_bk = f"bkup_vehicle_{safe}"
    log_bk = f"bkup_maint_log_{safe}"
    cur.execute(f"DROP TABLE IF EXISTS `{veh_bk}`")
    cur.execute(f"DROP TABLE IF EXISTS `{log_bk}`")
    cur.execute(f"CREATE TABLE `{veh_bk}` AS SELECT * FROM vehicle")
    cur.execute(f"CREATE TABLE `{log_bk}` AS SELECT * FROM vehicle_maintenance_log")
    return veh_bk, log_bk


def load_type_map(cur) -> dict[str, int]:
    cur.execute(
        "SELECT MAINT_CODE, VEHICLE_MAINTENANCE_TYPEID FROM vehicle_maintenance_type "
        "WHERE STATUS!=1"
    )
    return {r[0]: r[1] for r in cur.fetchall()}


def load_vin_map(cur) -> dict[str, int]:
    cur.execute(
        "SELECT UPPER(TRIM(VINNUMBER)), VEHICLEID FROM vehicle "
        "WHERE STATUS IN (0,4) AND VINNUMBER IS NOT NULL AND VINNUMBER!=''"
    )
    return {r[0]: r[1] for r in cur.fetchall() if r[0]}


def prefer_vehicle_number(existing: str | None, excel_van: str) -> str:
    excel = normalize_van(excel_van)
    if not excel:
        return (existing or "").strip()
    ex = (existing or "").strip()
    if not ex:
        return excel
    # Keep station-prefixed IDs like CDV-MVPG-17 when Excel only has CDV-17
    if excel.upper() in ex.upper() or ex.upper().endswith(excel.upper()):
        return ex
    return excel


def read_inventory(ws):
    rows = []
    for r in ws.iter_rows(min_row=3, max_row=ws.max_row, values_only=True):
        if not r or not r[1]:
            continue
        rows.append(
            {
                "van": r[0],
                "vin": str(r[1]).strip().upper(),
                "plate": r[2],
                "type": r[3],  # service tier
                "ownership": r[4],
                "provider": r[5],
                "year": r[6],
                "status": r[7],
                "reg_expiry": r[8],
                "reg_flag": r[9],
                "ground_cat": r[10],
                "issue": r[11],
                "last_oil": r[12],
                "next_pm": r[13],
                "recall": r[14],
                "tickets": r[15],
                "at_dealer": r[16],
                "support_case": r[17],
                "last_tire": r[18],
                "tire_note": r[19],
            }
        )
    return rows


def read_history(ws):
    rows = []
    for r in ws.iter_rows(min_row=5, max_row=ws.max_row, values_only=True):
        if not r or not r[0]:
            continue
        rows.append(
            {
                "van": r[0],
                "date": r[1],
                "source": r[2],
                "activity": r[3],
                "description": r[4],
                "outcome": r[5],
            }
        )
    return rows


def upsert_vehicles(cur, inv_rows, vin_map, type_map, dry_run: bool):
    inserted = updated = 0
    van_to_vin = {}
    for row in inv_rows:
        vin = row["vin"]
        van = normalize_van(str(row["van"] or ""))
        if van:
            van_to_vin[van.upper()] = vin
            # also index short aliases: CDV-17 from CDV-MVPG-17 handled via inventory van key
            van_to_vin[van.upper().replace("MVPG-", "")] = vin

        op, st = op_status(str(row["status"] or ""))
        ownership = (str(row["ownership"]).strip().upper() if row["ownership"] else "")
        vtype = ownership_to_vehicle_type(ownership)
        tier = (str(row["type"]).strip() if row["type"] else "")
        provider = (str(row["provider"]).strip() if row["provider"] else "")
        year = row["year"] if row["year"] not in (None, "") else None
        plate = (str(row["plate"]).strip() if row["plate"] else "")
        reason_code = (str(row["ground_cat"]).strip()[:30] if row["ground_cat"] else "")
        reason_msg = (str(row["issue"]).strip()[:100] if row["issue"] else "")
        out_repair = 1 if truthy(row["at_dealer"]) else 0
        reg_exp = sql_date(row["reg_expiry"])

        existing_id = vin_map.get(vin)
        if existing_id:
            cur.execute(
                "SELECT VEHICLENUMBER FROM vehicle WHERE VEHICLEID=%s", (existing_id,)
            )
            existing_num = cur.fetchone()
            existing_num = existing_num[0] if existing_num else ""
            vnum = prefer_vehicle_number(existing_num, str(row["van"] or ""))
            action = "UPDATE"
            sql = f"""
                UPDATE vehicle SET
                  VEHICLENUMBER={sql_str(vnum)},
                  LICENSEPLATE={sql_str(plate)},
                  VEHICLETYPE={sql_str(vtype)},
                  OWNERSHIPTYPE={sql_str(ownership)},
                  PROVIDER={sql_str(provider)},
                  SERVICETIER={sql_str(tier)},
                  VEHICLEYEAR={year if year is not None else 'NULL'},
                  OPERATIONALSTATUS={op},
                  STATUS={st},
                  STATUSREASONCODE={sql_str(reason_code)},
                  STATUSREASONMSG={sql_str(reason_msg)},
                  REGISTRATIONEXPIRY={reg_exp},
                  STATION={sql_str(STATION)},
                  OUT_FOR_REPAIR={out_repair},
                  UPDATE_USER={sql_str(BATCH_TAG)},
                  UPDATE_DATE=NOW()
                WHERE VEHICLEID={existing_id}
            """
            if not dry_run:
                cur.execute(sql)
                cur.execute(
                    "INSERT INTO fleet_import_touch (BATCH_TAG, VINNUMBER, VEHICLEID, ACTION, CREATE_DATE) "
                    "VALUES (%s,%s,%s,%s,NOW())",
                    (BATCH_TAG, vin, existing_id, action),
                )
            updated += 1
            vehicle_id = existing_id
        else:
            vnum = normalize_van(str(row["van"] or "")) or vin[-8:]
            action = "INSERT"
            sql = f"""
                INSERT INTO vehicle (
                  ENTITYID, VEHICLENUMBER, VINNUMBER, VEHICLETYPE, LICENSEPLATE,
                  REGISTRATIONEXPIRY, SERVICETIER, OPERATIONALSTATUS,
                  STATUSREASONCODE, STATUSREASONMSG, PROVIDER, VEHICLEYEAR,
                  OWNERSHIPTYPE, STATION, OUT_FOR_REPAIR,
                  CREATE_USER, CREATE_DATE, STATUS
                ) VALUES (
                  {ENTITY_ID}, {sql_str(vnum)}, {sql_str(vin)}, {sql_str(vtype)}, {sql_str(plate)},
                  {reg_exp}, {sql_str(tier)}, {op},
                  {sql_str(reason_code)}, {sql_str(reason_msg)}, {sql_str(provider)},
                  {year if year is not None else 'NULL'},
                  {sql_str(ownership)}, {sql_str(STATION)}, {out_repair},
                  {sql_str(BATCH_TAG)}, NOW(), {st}
                )
            """
            if not dry_run:
                cur.execute(sql)
                vehicle_id = cur.lastrowid
                vin_map[vin] = vehicle_id
                cur.execute(
                    "INSERT INTO fleet_import_touch (BATCH_TAG, VINNUMBER, VEHICLEID, ACTION, CREATE_DATE) "
                    "VALUES (%s,%s,%s,%s,NOW())",
                    (BATCH_TAG, vin, vehicle_id, action),
                )
            else:
                vehicle_id = -1
            inserted += 1

        # Optional oil PM snapshot as a maintenance row (skip if no real oil date)
        if row["last_oil"] is not None and sql_date(row["last_oil"]) != "NULL":
            oil_code = "OIL_CHANGE"
            tid = type_map.get(oil_code)
            desc = "Imported Last Oil Change from Fleet Inventory Report"
            next_pm = sql_date(row["next_pm"])
            if not dry_run and vehicle_id > 0:
                # de-dupe same batch + same service date + oil
                cur.execute(
                    """
                    SELECT MAINT_LOGID FROM vehicle_maintenance_log
                    WHERE VEHICLEID=%s AND CREATE_USER=%s AND MAINT_CODE=%s
                      AND DATE(SERVICE_DATE)=DATE(%s) AND STATUS=0
                    LIMIT 1
                    """,
                    (vehicle_id, BATCH_TAG, oil_code, row["last_oil"]),
                )
                if not cur.fetchone():
                    cur.execute(
                        f"""
                        INSERT INTO vehicle_maintenance_log (
                          VEHICLEID, ENTITYID, MAINT_TYPEID, MAINT_CODE, MAINT_CATEGORY,
                          SERVICE_DATE, COMPLETED_DATE, DESCRIPTION, IS_OPEN,
                          NEXT_SERVICE_DATE, CREATE_USER, CREATE_DATE, STATUS
                        ) VALUES (
                          {vehicle_id}, {ENTITY_ID},
                          {tid if tid else 'NULL'}, {sql_str(oil_code)}, 'PM',
                          {sql_date(row['last_oil'])}, {sql_date(row['last_oil'])},
                          {sql_str(desc)}, 0, {next_pm},
                          {sql_str(BATCH_TAG)}, NOW(), 0
                        )
                        """
                    )

    return inserted, updated, van_to_vin


def insert_history(cur, hist_rows, vin_map, van_to_vin, type_map, dry_run: bool):
    inserted = skipped = 0
    for row in hist_rows:
        van = normalize_van(str(row["van"] or ""))
        vin = van_to_vin.get(van.upper()) if van else None
        if not vin:
            # try alias CDV-MVPG-17 <-> CDV-17
            alt = van.upper().replace("MVPG-", "") if van else ""
            vin = van_to_vin.get(alt)
        if not vin:
            skipped += 1
            continue
        vehicle_id = vin_map.get(vin)
        if not vehicle_id:
            skipped += 1
            continue

        code = map_activity(str(row["activity"] or ""))
        tid = type_map.get(code)
        svc = sql_date(row["date"])
        source = str(row["source"] or "").strip()
        activity = str(row["activity"] or "").strip()
        description = str(row["description"] or "").strip()
        outcome = str(row["outcome"] or "").strip()
        parts = []
        if activity:
            parts.append(activity)
        if description:
            parts.append(description)
        if outcome:
            parts.append(f"Outcome: {outcome}")
        if source:
            parts.append(f"Source: {source}")
        full_desc = " | ".join(parts)[:2000]
        shop = source[:150] if source else ""

        # open if outcome suggests still out
        outcome_l = outcome.lower()
        is_open = 1 if any(x in outcome_l for x in ("await", "backorder", "at dealer", "open")) else 0
        completed = "NULL" if is_open else (svc if svc != "NULL" else "NULL")

        if dry_run:
            inserted += 1
            continue

        # de-dupe within batch: same vehicle + date + activity + description hash
        fingerprint = hashlib.sha1(
            f"{vehicle_id}|{svc}|{activity}|{description}|{outcome}".encode("utf-8", "ignore")
        ).hexdigest()[:16]
        cur.execute(
            """
            SELECT MAINT_LOGID FROM vehicle_maintenance_log
            WHERE VEHICLEID=%s AND CREATE_USER=%s AND STATUS=0
              AND DESCRIPTION LIKE %s
            LIMIT 1
            """,
            (vehicle_id, BATCH_TAG, f"%[{fingerprint}]%"),
        )
        if cur.fetchone():
            skipped += 1
            continue

        desc_tagged = (full_desc + f" [{fingerprint}]")[:2000]
        cur.execute(
            f"""
            INSERT INTO vehicle_maintenance_log (
              VEHICLEID, ENTITYID, MAINT_TYPEID, MAINT_CODE, MAINT_CATEGORY,
              SERVICE_DATE, COMPLETED_DATE, SHOP_VENDOR, DESCRIPTION,
              IS_OPEN, CREATE_USER, CREATE_DATE, STATUS
            ) VALUES (
              {vehicle_id}, {ENTITY_ID},
              {tid if tid else 'NULL'}, {sql_str(code)}, 'IMPORT',
              {svc}, {completed}, {sql_str(shop)}, {sql_str(desc_tagged)},
              {is_open}, {sql_str(BATCH_TAG)}, NOW(), 0
            )
            """
        )
        inserted += 1
    return inserted, skipped


def main():
    global BATCH_TAG
    ap = argparse.ArgumentParser()
    ap.add_argument("--excel", type=Path, default=DEFAULT_EXCEL)
    ap.add_argument("--dry-run", action="store_true")
    ap.add_argument("--batch-tag", default=None)
    args = ap.parse_args()
    if args.batch_tag:
        BATCH_TAG = args.batch_tag

    if not args.excel.exists():
        print(f"Excel not found: {args.excel}")
        sys.exit(1)

    sha1 = hashlib.sha1(args.excel.read_bytes()).hexdigest()
    wb = openpyxl.load_workbook(args.excel, data_only=True)
    inv = read_inventory(wb["Current Issues (Full Inventory)"])
    hist = read_history(wb["Vehicle History"])
    print(f"BATCH_TAG={BATCH_TAG}")
    print(f"Inventory rows={len(inv)}  History rows={len(hist)}  dry_run={args.dry_run}")

    cn = connect()
    cur = cn.cursor()
    try:
        ensure_batch_tables(cur)
        if not args.dry_run:
            veh_bk, log_bk = backup_tables(cur, BATCH_TAG)
            print(f"Backups: {veh_bk}, {log_bk}")
            cur.execute(
                "INSERT INTO fleet_import_batch (BATCH_TAG, EXCEL_PATH, FILE_SHA1, STARTED_AT) "
                "VALUES (%s,%s,%s,NOW())",
                (BATCH_TAG, str(args.excel), sha1),
            )
        else:
            veh_bk = log_bk = None

        type_map = load_type_map(cur)
        vin_map = load_vin_map(cur)
        v_ins, v_upd, van_to_vin = upsert_vehicles(cur, inv, vin_map, type_map, args.dry_run)
        # refresh vin map after inserts
        if not args.dry_run:
            vin_map = load_vin_map(cur)
        h_ins, h_skip = insert_history(cur, hist, vin_map, van_to_vin, type_map, args.dry_run)

        if not args.dry_run:
            cur.execute(
                """
                UPDATE fleet_import_batch SET
                  FINISHED_AT=NOW(),
                  VEHICLES_INSERTED=%s,
                  VEHICLES_UPDATED=%s,
                  HISTORY_INSERTED=%s,
                  NOTES=%s
                WHERE BATCH_TAG=%s
                """,
                (
                    v_ins,
                    v_upd,
                    h_ins,
                    f"history_skipped={h_skip}; backups={veh_bk},{log_bk}",
                    BATCH_TAG,
                ),
            )
            cn.commit()
        else:
            cn.rollback()

        print(f"Vehicles INSERT={v_ins} UPDATE={v_upd}")
        print(f"History INSERT={h_ins} SKIP={h_skip}")
        print("Done.")
        if not args.dry_run:
            print(f"Rollback: mysql ... < rollback_fleet_inventory.sql  (set @batch:='{BATCH_TAG}')")
    except Exception:
        cn.rollback()
        raise
    finally:
        cur.close()
        cn.close()


if __name__ == "__main__":
    main()
